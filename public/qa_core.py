"""
QA Call Randomizer - core engine.

Pure data logic: loading, validation, duration normalisation, sampling and
Excel output.  Deliberately free of any GUI import so it can be unit-tested
and reused head-less.

All core field identification is done by FIXED EXCEL COLUMN POSITION,
never by header name.
"""

from __future__ import annotations

import datetime as _dt
import math
import os
import random
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Sequence, Tuple

import pandas as pd

APP_NAME = "QA Call Randomizer"
APP_VERSION = "1.0.0"

# --------------------------------------------------------------------------
# Fixed column mapping (Excel letter -> 0-based position)
# --------------------------------------------------------------------------
FIELD_COLUMNS: List[Tuple[str, str, int]] = [
    # (internal field, excel letter, 0-based index)
    ("Transaction", "A", 0),
    ("Agent", "C", 2),
    ("Recording Type", "D", 3),
    ("Date & Time", "J", 9),
    ("Duration", "L", 11),
    ("Login ID", "AH", 33),
    ("GenConnID", "AN", 39),
    ("Post_Route_Data", "AO", 40),
]

MIN_REQUIRED_COLUMNS = 41  # through Column AO
VOICE_AND_SCREEN = "voice and screen"

OUTPUT_COLUMNS = [
    "Transaction",
    "Agent",
    "Recording Type",
    "Date & Time",
    "Duration",
    "Login ID",
    "GenConnID",
    "Post_Route_Data",
    "Sample Category",
    "Recording Priority",
    "Code Match",
    "Selection Reason",
    "Duration (Normalized)",
]


class QaRandomizerError(Exception):
    """User-facing error with a friendly message."""


# --------------------------------------------------------------------------
# Configuration
# --------------------------------------------------------------------------
@dataclass
class Bucket:
    label: str
    min_seconds: int
    max_seconds: Optional[int]  # None = open ended
    target: int

    def contains(self, seconds: pd.Series) -> pd.Series:
        mask = seconds >= self.min_seconds
        if self.max_seconds is not None:
            mask &= seconds <= self.max_seconds
        return mask

    @property
    def short_label(self) -> str:
        return self.label.replace(" Minutes", "").replace("+ ", "+")


def default_buckets() -> List[Bucket]:
    """5-20 min x10, 21-35 min x5, 36+ min x5 (note the intentional gaps)."""
    return [
        Bucket("5\u201320 Minutes", 300, 1200, 10),
        Bucket("21\u201335 Minutes", 1260, 2100, 5),
        Bucket("36+ Minutes", 2160, None, 5),
    ]


@dataclass
class SamplingConfig:
    buckets: List[Bucket] = field(default_factory=default_buckets)
    prioritize_voice_and_screen: bool = True
    code_filter_enabled: bool = False
    codes: List[str] = field(default_factory=list)
    match_mode: str = "contains"        # "contains" | "exact"
    multi_code_logic: str = "any"       # "any" | "all"
    seed: Optional[int] = None
    numeric_duration_units: str = "auto"  # auto | day_fraction | minutes | seconds

    @property
    def calls_per_agent(self) -> int:
        return sum(b.target for b in self.buckets)

    def code_filter_label(self) -> str:
        if not self.code_filter_enabled:
            return "OFF"
        n = len(self.codes)
        return f"ON \u2014 {n} code{'s' if n != 1 else ''}"


# --------------------------------------------------------------------------
# Duration normalisation
# --------------------------------------------------------------------------
_HMS_RE = re.compile(r"^\d{1,4}(:\d{1,2}){1,2}(\.\d+)?$")
_UNIT_RE = re.compile(
    r"^(?:(?P<h>\d+(?:\.\d+)?)\s*h(?:r|rs|our|ours)?)?\s*"
    r"(?:(?P<m>\d+(?:\.\d+)?)\s*m(?:in|ins|inute|inutes)?)?\s*"
    r"(?:(?P<s>\d+(?:\.\d+)?)\s*s(?:ec|ecs|econd|econds)?)?$",
    re.IGNORECASE,
)

# Auto heuristic for bare numbers:
#   < 1        -> Excel fractional day  (0.0125 -> 18:00)
#   1 .. 239   -> minutes               (41.33  -> 41:20)
#   >= 240     -> seconds               (2480   -> 41:20)
_AUTO_MINUTES_CEILING = 240


def _numeric_to_seconds(value: float, units: str) -> Optional[float]:
    if value != value or value in (float("inf"), float("-inf")):  # NaN / inf
        return None
    if value < 0:
        return None
    if units == "day_fraction":
        return value * 86400.0
    if units == "minutes":
        return value * 60.0
    if units == "seconds":
        return value
    # auto
    if value < 1:
        return value * 86400.0
    if value < _AUTO_MINUTES_CEILING:
        return value * 60.0
    return value


def parse_duration(value: Any, numeric_units: str = "auto") -> Optional[int]:
    """Normalise any reasonable duration representation to whole seconds.

    Supports: datetime.time, datetime.timedelta / pandas Timedelta,
    datetime.datetime (Excel 1899-12-30 base), Excel fractional days,
    "HH:MM:SS", "MM:SS", "1h 05m 32s" and bare numeric values.
    Returns None when the value cannot be understood.
    """
    if value is None:
        return None
    if isinstance(value, bool):
        return None

    # pandas / numpy NaN
    try:
        if value is pd.NaT or (isinstance(value, float) and math.isnan(value)):
            return None
    except Exception:  # pragma: no cover - defensive
        pass

    if isinstance(value, _dt.timedelta):
        secs = value.total_seconds()
        return int(round(secs)) if secs >= 0 else None
    if isinstance(value, pd.Timedelta):  # pragma: no cover - subclass of timedelta
        return int(round(value.total_seconds()))
    if isinstance(value, _dt.datetime):
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, _dt.time):
        return value.hour * 3600 + value.minute * 60 + value.second

    if isinstance(value, (int, float)):
        secs = _numeric_to_seconds(float(value), numeric_units)
        return int(round(secs)) if secs is not None else None

    text = str(value).strip()
    if not text or text.lower() in {"nan", "nat", "none", "null", "-", "--"}:
        return None

    negative = text.startswith("-")
    if negative:
        return None

    # "1 day, 0:12:33"
    day_part = 0
    m = re.match(r"^(\d+)\s*days?[,\s]+(.*)$", text, re.IGNORECASE)
    if m:
        day_part = int(m.group(1)) * 86400
        text = m.group(2).strip()

    if ":" in text and _HMS_RE.match(text):
        parts = text.split(":")
        try:
            nums = [float(p) for p in parts]
        except ValueError:
            return None
        if len(parts) == 2:          # MM:SS  (05:32 -> 332 s)
            secs = nums[0] * 60 + nums[1]
        else:                        # HH:MM:SS
            secs = nums[0] * 3600 + nums[1] * 60 + nums[2]
        return int(round(secs + day_part))

    # bare number, possibly with thousands separators
    cleaned = text.replace(",", "")
    try:
        numeric = float(cleaned)
    except ValueError:
        numeric = None
    if numeric is not None:
        secs = _numeric_to_seconds(numeric, numeric_units)
        return int(round(secs + day_part)) if secs is not None else None

    um = _UNIT_RE.match(text)
    if um and any(um.group(g) for g in ("h", "m", "s")):
        secs = (
            float(um.group("h") or 0) * 3600
            + float(um.group("m") or 0) * 60
            + float(um.group("s") or 0)
        )
        return int(round(secs + day_part))

    return None


def format_hhmmss(seconds: Optional[int]) -> str:
    if seconds is None or (isinstance(seconds, float) and math.isnan(seconds)):
        return ""
    seconds = int(seconds)
    h, rem = divmod(seconds, 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


# --------------------------------------------------------------------------
# Loading & preparation
# --------------------------------------------------------------------------
def _norm_text(value: Any) -> str:
    if value is None or value is pd.NaT:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def load_source(path: str | os.PathLike) -> pd.DataFrame:
    """Read the source workbook/CSV exactly as provided (read-only)."""
    p = Path(path)
    if not p.exists():
        raise QaRandomizerError("The selected file could not be found.")
    suffix = p.suffix.lower()
    try:
        if suffix in {".csv", ".txt", ".tsv"}:
            sep = "\t" if suffix == ".tsv" else ","
            try:
                df = pd.read_csv(p, header=0, dtype=object, sep=sep,
                                 keep_default_na=False, na_filter=False,
                                 engine="python")
            except UnicodeDecodeError:
                df = pd.read_csv(p, header=0, dtype=object, sep=sep,
                                 keep_default_na=False, na_filter=False,
                                 engine="python", encoding="latin-1")
        elif suffix == ".xls":
            df = pd.read_excel(p, header=0, dtype=object)
        elif suffix in {".xlsx", ".xlsm"}:
            df = pd.read_excel(p, header=0, dtype=object, engine="openpyxl")
        else:
            raise QaRandomizerError(
                "Unable to read this file. Please upload a valid Excel or CSV file "
                "(.xlsx, .xls or .csv)."
            )
    except QaRandomizerError:
        raise
    except ImportError as exc:
        raise QaRandomizerError(
            "This .xls file needs the legacy Excel reader. Please re-save it as .xlsx "
            "and try again."
        ) from exc
    except Exception as exc:
        raise QaRandomizerError(
            "Unable to read this file. Please upload a valid Excel or CSV file."
        ) from exc

    if df.shape[1] < MIN_REQUIRED_COLUMNS:
        raise QaRandomizerError(
            "The uploaded file does not contain the required columns through Column AO.\n\n"
            f"Columns found: {df.shape[1]}. At least {MIN_REQUIRED_COLUMNS} are required."
        )
    if len(df) == 0:
        raise QaRandomizerError("The uploaded file contains no call records.")
    return df


def detected_headers(df: pd.DataFrame) -> List[Tuple[str, str, str]]:
    """[(excel letter, internal field, actual header text in the file)]"""
    out = []
    for field_name, letter, idx in FIELD_COLUMNS:
        header = str(df.columns[idx]) if idx < df.shape[1] else ""
        if header.startswith("Unnamed:"):
            header = "(blank header)"
        out.append((letter, field_name, header))
    return out


@dataclass
class ValidationReport:
    total_rows: int = 0
    valid_agent: int = 0
    missing_agent: int = 0
    valid_duration: int = 0
    invalid_duration: int = 0
    voice_and_screen: int = 0
    other_recording: int = 0
    post_route_populated: int = 0
    usable_rows: int = 0
    duplicate_ids: int = 0

    def as_rows(self) -> List[Tuple[str, Any]]:
        return [
            ("Total rows", self.total_rows),
            ("Valid agent", self.valid_agent),
            ("Missing agent", self.missing_agent),
            ("Valid duration", self.valid_duration),
            ("Invalid duration", self.invalid_duration),
            ("Voice And Screen", self.voice_and_screen),
            ("Other recording types", self.other_recording),
            ("Post_Route_Data populated", self.post_route_populated),
            ("Rows usable for sampling", self.usable_rows),
        ]


@dataclass
class PreparedData:
    frame: pd.DataFrame               # working frame (normalised helper columns)
    source: pd.DataFrame              # untouched source rows
    validation: ValidationReport
    headers: List[Tuple[str, str, str]]
    source_path: str
    agents: List[str] = field(default_factory=list)
    date_range: Optional[Tuple[Any, Any]] = None


def prepare_data(df: pd.DataFrame, source_path: str,
                 numeric_units: str = "auto") -> PreparedData:
    """Build the normalised working frame + validation report."""
    work = pd.DataFrame(index=df.index)
    for field_name, _letter, idx in FIELD_COLUMNS:
        work[field_name] = df.iloc[:, idx].values

    work["_agent_display"] = work["Agent"].map(_norm_text)
    work["_agent_key"] = work["_agent_display"].str.casefold()

    # Duration: normalise on unique values only (fast on large files)
    dur_series = work["Duration"]
    try:
        uniques = pd.unique(dur_series)
        lookup = {}
        for val in uniques:
            try:
                lookup[val] = parse_duration(val, numeric_units)
            except TypeError:  # unhashable
                pass
        if len(lookup) == len(uniques):
            work["_duration_seconds"] = dur_series.map(
                lambda v: lookup.get(v, parse_duration(v, numeric_units))
            )
        else:
            raise TypeError
    except TypeError:
        work["_duration_seconds"] = dur_series.map(
            lambda v: parse_duration(v, numeric_units)
        )
    work["_duration_seconds"] = pd.to_numeric(work["_duration_seconds"], errors="coerce")
    work["_duration_text"] = work["_duration_seconds"].map(
        lambda s: format_hhmmss(int(s)) if pd.notna(s) else ""
    )

    rec = work["Recording Type"].map(_norm_text).str.casefold()
    work["_is_vas"] = rec.eq(VOICE_AND_SCREEN)

    work["_post_route"] = work["Post_Route_Data"].map(_norm_text)

    # Unique identifier: GenConnID -> Transaction -> internal row id
    gen = work["GenConnID"].map(_norm_text)
    txn = work["Transaction"].map(_norm_text)
    uid = gen.where(gen != "", txn)
    fallback = pd.Series([f"__ROW__{i}" for i in range(len(work))], index=work.index)
    work["_uid"] = uid.where(uid != "", fallback)

    has_agent = work["_agent_key"] != ""
    has_duration = work["_duration_seconds"].notna()

    report = ValidationReport(
        total_rows=int(len(work)),
        valid_agent=int(has_agent.sum()),
        missing_agent=int((~has_agent).sum()),
        valid_duration=int(has_duration.sum()),
        invalid_duration=int((~has_duration).sum()),
        voice_and_screen=int(work["_is_vas"].sum()),
        other_recording=int((~work["_is_vas"]).sum()),
        post_route_populated=int((work["_post_route"] != "").sum()),
        usable_rows=int((has_agent & has_duration).sum()),
        duplicate_ids=int(work["_uid"].duplicated().sum()),
    )

    agents = sorted(
        work.loc[has_agent, "_agent_display"].unique().tolist(),
        key=lambda s: s.casefold(),
    )

    date_range = None
    try:
        dts = pd.to_datetime(work["Date & Time"], errors="coerce")
        if dts.notna().any():
            date_range = (dts.min(), dts.max())
    except Exception:  # pragma: no cover - defensive
        date_range = None

    return PreparedData(
        frame=work,
        source=df,
        validation=report,
        headers=detected_headers(df),
        source_path=str(source_path),
        agents=agents,
        date_range=date_range,
    )


# --------------------------------------------------------------------------
# Post_Route_Data code matching
# --------------------------------------------------------------------------
def parse_codes(text: str) -> List[str]:
    if not text:
        return []
    parts = re.split(r"[,;\n\r\t]+", text)
    return [p.strip() for p in parts if p.strip()]


def build_code_mask(post_route: pd.Series, config: SamplingConfig) -> pd.Series:
    """Vectorised case-insensitive matching against Post_Route_Data."""
    codes = [c.casefold() for c in config.codes]
    if not codes:
        return pd.Series(True, index=post_route.index)
    values = post_route.str.casefold()
    masks = []
    for code in codes:
        if config.match_mode == "exact":
            masks.append(values.eq(code))
        else:
            masks.append(values.str.contains(re.escape(code), regex=True, na=False))
    result = masks[0].copy()
    for m in masks[1:]:
        result = (result & m) if config.multi_code_logic == "all" else (result | m)
    return result


# --------------------------------------------------------------------------
# Sampling
# --------------------------------------------------------------------------
REASON_VAS = "Voice And Screen priority"
REASON_FALLBACK = "Voice And Screen priority + random fallback"
REASON_RANDOM = "Random selection"

STATUS_COMPLETE = "Complete"
STATUS_PARTIAL = "Partial \u2014 Insufficient Calls"
STATUS_NONE = "No Eligible Calls"


@dataclass
class SamplingResult:
    selected: pd.DataFrame
    summary: pd.DataFrame
    config: SamplingConfig
    seed_used: int
    validation: ValidationReport
    source_path: str
    agents_processed: int
    calls_analyzed: int
    expected_total: int
    warnings: List[str] = field(default_factory=list)
    run_timestamp: _dt.datetime = field(default_factory=_dt.datetime.now)

    @property
    def calls_selected(self) -> int:
        return int(len(self.selected))

    @property
    def voice_and_screen_selected(self) -> int:
        return int((self.selected["Recording Priority"] == "Voice And Screen").sum()) \
            if len(self.selected) else 0

    @property
    def other_selected(self) -> int:
        return self.calls_selected - self.voice_and_screen_selected

    @property
    def completion_pct(self) -> float:
        if not self.expected_total:
            return 0.0
        return round(100.0 * self.calls_selected / self.expected_total, 1)


def run_sampling(prepared: PreparedData, config: SamplingConfig,
                 progress: Optional[Callable[[str, int], None]] = None) -> SamplingResult:
    """Execute the QA sampling algorithm. Never mutates the source frame."""

    def report(message: str, pct: int) -> None:
        if progress:
            progress(message, pct)

    report("Analyzing calls\u2026", 5)
    work = prepared.frame
    seed_used = config.seed if config.seed is not None else random.SystemRandom().randrange(1, 10**9)
    rng = random.Random(seed_used)

    base_valid = work[(work["_agent_key"] != "") & (work["_duration_seconds"].notna())]

    report("Grouping agents\u2026", 12)
    # The agent universe comes from all valid rows, BEFORE the code filter, so an
    # agent whose calls are all filtered out is still reported (No Eligible Calls).
    agent_names = (base_valid.groupby("_agent_key")["_agent_display"].first().to_dict())
    agent_keys = sorted(agent_names.keys())

    report("Applying Post_Route_Data filters\u2026", 20)
    eligible = base_valid
    if config.code_filter_enabled and config.codes:
        code_mask = build_code_mask(eligible["_post_route"], config)
        eligible = eligible[code_mask]

    grouped = {key: idx for key, idx in eligible.groupby("_agent_key").groups.items()}

    selected_uids: set = set()
    picked_rows: List[Dict[str, Any]] = []
    summary_rows: List[Dict[str, Any]] = []
    warnings: List[str] = []

    total_agents = max(len(agent_keys), 1)
    for i, key in enumerate(agent_keys):
        agent_frame = eligible.loc[grouped[key]] if key in grouped else eligible.iloc[0:0]
        display_name = agent_names[key]
        pct = 25 + int(65 * (i / total_agents))
        if i % 5 == 0 or total_agents < 20:
            report(f"Randomizing calls\u2026 ({i + 1}/{total_agents} agents)", pct)

        row_summary: Dict[str, Any] = {"Agent": display_name}
        agent_selected = 0
        agent_target = 0
        agent_vas = 0

        durations = agent_frame["_duration_seconds"]
        for bucket in config.buckets:
            pool = agent_frame[bucket.contains(durations)]
            # never reuse a call already selected in another bucket / duplicate ids
            if selected_uids:
                pool = pool[~pool["_uid"].isin(selected_uids)]
            pool = pool[~pool["_uid"].duplicated(keep="first")]

            target = bucket.target
            agent_target += target

            vas_idx = pool.index[pool["_is_vas"]].tolist()
            other_idx = pool.index[~pool["_is_vas"]].tolist()

            chosen: List[Tuple[Any, str]] = []
            if config.prioritize_voice_and_screen:
                take_vas = min(target, len(vas_idx))
                for idx in rng.sample(vas_idx, take_vas):
                    chosen.append((idx, REASON_VAS))
                remaining = target - take_vas
                if remaining > 0 and other_idx:
                    take_other = min(remaining, len(other_idx))
                    reason = REASON_FALLBACK if take_vas > 0 else REASON_RANDOM
                    for idx in rng.sample(other_idx, take_other):
                        chosen.append((idx, reason))
            else:
                all_idx = vas_idx + other_idx
                take = min(target, len(all_idx))
                for idx in rng.sample(all_idx, take):
                    chosen.append((idx, REASON_RANDOM))

            for idx, reason in chosen:
                row = work.loc[idx]
                selected_uids.add(row["_uid"])
                is_vas = bool(row["_is_vas"])
                agent_vas += 1 if is_vas else 0
                picked_rows.append(
                    {
                        "Transaction": row["Transaction"],
                        "Agent": row["Agent"],
                        "Recording Type": row["Recording Type"],
                        "Date & Time": row["Date & Time"],
                        "Duration": row["Duration"],
                        "Login ID": row["Login ID"],
                        "GenConnID": row["GenConnID"],
                        "Post_Route_Data": row["Post_Route_Data"],
                        "Sample Category": bucket.label,
                        "Recording Priority": "Voice And Screen" if is_vas else "Other Recording Type",
                        "Code Match": "Matched" if (config.code_filter_enabled and config.codes) else "Not Applicable",
                        "Selection Reason": reason,
                        "Duration (Normalized)": row["_duration_text"],
                        "_agent_key": key,
                        "_duration_seconds": row["_duration_seconds"],
                    }
                )

            row_summary[f"{bucket.label} Target"] = target
            row_summary[f"{bucket.label} Selected"] = len(chosen)
            row_summary[f"{bucket.label} Available"] = len(pool)
            agent_selected += len(chosen)

        row_summary["Total Target"] = agent_target
        row_summary["Total Selected"] = agent_selected
        row_summary["Voice And Screen Selected"] = agent_vas
        row_summary["Other Recording Types Selected"] = agent_selected - agent_vas
        row_summary["Code Filter"] = config.code_filter_label()
        if agent_selected == 0:
            row_summary["Status"] = STATUS_NONE
        elif agent_selected < agent_target:
            row_summary["Status"] = STATUS_PARTIAL
        else:
            row_summary["Status"] = STATUS_COMPLETE

        if agent_selected < agent_target:
            detail = ", ".join(
                f"{b.label}: {row_summary[f'{b.label} Selected']}/{b.target}"
                for b in config.buckets
            )
            warnings.append(
                f"{display_name} \u2014 {agent_selected}/{agent_target} selected ({detail})"
            )
        summary_rows.append(row_summary)

    report("Generating Excel file\u2026", 92)

    selected_df = pd.DataFrame(picked_rows, columns=OUTPUT_COLUMNS + ["_agent_key", "_duration_seconds"]) \
        if picked_rows else pd.DataFrame(columns=OUTPUT_COLUMNS + ["_agent_key", "_duration_seconds"])
    if len(selected_df):
        selected_df = selected_df.sort_values(
            ["_agent_key", "Sample Category", "_duration_seconds"], kind="mergesort"
        ).reset_index(drop=True)

    summary_cols = ["Agent"]
    for b in config.buckets:
        summary_cols += [f"{b.label} Target", f"{b.label} Selected"]
    summary_cols += [
        "Total Target", "Total Selected",
        "Voice And Screen Selected", "Other Recording Types Selected",
        "Code Filter", "Status",
    ]
    summary_df = pd.DataFrame(summary_rows, columns=summary_cols) if summary_rows \
        else pd.DataFrame(columns=summary_cols)

    result = SamplingResult(
        selected=selected_df,
        summary=summary_df,
        config=config,
        seed_used=seed_used,
        validation=prepared.validation,
        source_path=prepared.source_path,
        agents_processed=len(agent_keys),
        calls_analyzed=int(len(work)),
        expected_total=len(agent_keys) * config.calls_per_agent,
        warnings=warnings,
    )
    report("Sample generated successfully.", 100)
    return result


# --------------------------------------------------------------------------
# Excel output
# --------------------------------------------------------------------------
NAVY = "0F2438"
STEEL = "3E6B8E"
GOLD = "C9A227"
LIGHT = "EEF2F6"
GREEN_FILL = "E4F1E4"
AMBER_FILL = "FDF3D8"
RED_FILL = "FBE3E3"


def _clean_cell(value: Any) -> Any:
    """Excel-safe value: NaN/NaT become blank, everything else is preserved as-is."""
    if value is None or value is pd.NaT:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, (pd.Timestamp,)):
        return value.to_pydatetime()
    return value


def _cell_text(value: Any) -> str:
    cleaned = _clean_cell(value)
    return "" if cleaned is None else str(cleaned)


def unique_output_path(path: Path) -> Path:
    """Never overwrite: append -1, -2 ... if needed."""
    if not path.exists():
        return path
    stem, suffix, parent = path.stem, path.suffix, path.parent
    n = 1
    while True:
        candidate = parent / f"{stem}-{n}{suffix}"
        if not candidate.exists():
            return candidate
        n += 1


def default_output_name(when: Optional[_dt.datetime] = None) -> str:
    when = when or _dt.datetime.now()
    return f"QA_Random_Call_Sample_{when:%Y-%m-%d_%H%M%S}.xlsx"


def write_output(result: SamplingResult, out_path: str | os.PathLike) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    path = unique_output_path(Path(out_path))
    wb = Workbook()

    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=NAVY)
    body_font = Font(name="Arial", size=10)
    title_font = Font(name="Arial", size=12, bold=True, color=NAVY)
    thin = Side(style="thin", color="D2DAE2")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style_header(ws, row: int, ncols: int) -> None:
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        ws.row_dimensions[row].height = 30

    def autosize(ws, df: pd.DataFrame, start_col: int = 1, min_w: int = 10, max_w: int = 42) -> None:
        for i, col in enumerate(df.columns):
            values = [_cell_text(v) for v in df[col].head(400).tolist()] + [str(col)]
            width = max(len(v) for v in values) + 3
            ws.column_dimensions[get_column_letter(start_col + i)].width = max(min_w, min(max_w, width))

    # ---------------- Sheet 1: Selected Calls ----------------
    ws1 = wb.active
    ws1.title = "Selected Calls"
    sel = result.selected[OUTPUT_COLUMNS].copy()
    ws1.append(OUTPUT_COLUMNS)
    for row in sel.itertuples(index=False):
        ws1.append([_clean_cell(v) for v in row])
    style_header(ws1, 1, len(OUTPUT_COLUMNS))
    for r in range(2, ws1.max_row + 1):
        for c in range(1, len(OUTPUT_COLUMNS) + 1):
            cell = ws1.cell(row=r, column=c)
            cell.font = body_font
            cell.border = border
        if r % 2 == 0:
            for c in range(1, len(OUTPUT_COLUMNS) + 1):
                ws1.cell(row=r, column=c).fill = PatternFill("solid", fgColor=LIGHT)
    ws1.freeze_panes = "A2"
    if ws1.max_row >= 1:
        ws1.auto_filter.ref = f"A1:{get_column_letter(len(OUTPUT_COLUMNS))}{max(ws1.max_row, 1)}"
    autosize(ws1, sel)

    # ---------------- Sheet 2: Sampling Summary ----------------
    ws2 = wb.create_sheet("Sampling Summary")
    summary = result.summary.copy()
    ws2.append(list(summary.columns))
    for row in summary.itertuples(index=False):
        ws2.append([_clean_cell(v) for v in row])
    style_header(ws2, 1, len(summary.columns))
    status_col = list(summary.columns).index("Status") + 1
    for r in range(2, ws2.max_row + 1):
        for c in range(1, len(summary.columns) + 1):
            cell = ws2.cell(row=r, column=c)
            cell.font = body_font
            cell.border = border
        status = ws2.cell(row=r, column=status_col).value
        fill = GREEN_FILL if status == STATUS_COMPLETE else (
            AMBER_FILL if status == STATUS_PARTIAL else RED_FILL)
        ws2.cell(row=r, column=status_col).fill = PatternFill("solid", fgColor=fill)

    # Totals row (live SUM formulas)
    if len(summary):
        total_row = ws2.max_row + 1
        ws2.cell(row=total_row, column=1, value="TOTAL").font = Font(
            name="Arial", size=10, bold=True, color=NAVY)
        for c, col_name in enumerate(summary.columns, start=1):
            if c == 1 or col_name in {"Status", "Code Filter"}:
                continue
            letter = get_column_letter(c)
            cell = ws2.cell(row=total_row, column=c,
                            value=f"=SUM({letter}2:{letter}{total_row - 1})")
            cell.font = Font(name="Arial", size=10, bold=True)
            cell.fill = PatternFill("solid", fgColor="DCE4EC")
            cell.border = border
        ws2.cell(row=total_row, column=1).fill = PatternFill("solid", fgColor="DCE4EC")
        ws2.cell(row=total_row, column=1).border = border
    ws2.freeze_panes = "A2"
    autosize(ws2, summary, min_w=12)

    # ---------------- Sheet 3: Run Details ----------------
    ws3 = wb.create_sheet("Run Details")
    cfg = result.config
    dist = " / ".join(f"{b.label}: {b.target}" for b in cfg.buckets)
    details = [
        ("Run date/time", result.run_timestamp.strftime("%Y-%m-%d %H:%M:%S")),
        ("Source filename", Path(result.source_path).name),
        ("Number of source records", result.calls_analyzed),
        ("Number of agents", result.agents_processed),
        ("Number of selected calls", result.calls_selected),
        ("Expected calls", result.expected_total),
        ("Completion", f"{result.completion_pct}%"),
        ("Requested calls per agent", cfg.calls_per_agent),
        ("Duration distribution", dist),
        ("Random seed", result.seed_used),
        ("Voice And Screen priority", "ON" if cfg.prioritize_voice_and_screen else "OFF"),
        ("Post_Route_Data filter status", cfg.code_filter_label()),
        ("Codes used", ", ".join(cfg.codes) if cfg.codes else "(none)"),
        ("Match mode", cfg.match_mode.title() if cfg.code_filter_enabled else "Not applied"),
        ("Multiple code logic",
         ("Any Code" if cfg.multi_code_logic == "any" else "All Codes")
         if cfg.code_filter_enabled else "Not applied"),
        ("Numeric duration units", cfg.numeric_duration_units),
        ("Voice And Screen selected", result.voice_and_screen_selected),
        ("Other recording types selected", result.other_selected),
        ("Rows with missing agent (excluded)", result.validation.missing_agent),
        ("Rows with invalid duration (excluded)", result.validation.invalid_duration),
        ("Application version", f"{APP_NAME} v{APP_VERSION}"),
    ]
    ws3.cell(row=1, column=1, value="Run Details").font = title_font
    ws3.append([])
    ws3.append(["Item", "Value"])
    style_header(ws3, 3, 2)
    for label, value in details:
        ws3.append([label, value])
    for r in range(4, ws3.max_row + 1):
        ws3.cell(row=r, column=1).font = Font(name="Arial", size=10, bold=True)
        ws3.cell(row=r, column=2).font = body_font
        ws3.cell(row=r, column=1).border = border
        ws3.cell(row=r, column=2).border = border
    ws3.column_dimensions["A"].width = 34
    ws3.column_dimensions["B"].width = 60

    wb.save(path)
    return path
